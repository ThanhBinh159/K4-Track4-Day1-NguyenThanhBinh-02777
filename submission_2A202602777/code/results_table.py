"""results_table.py — Lưu kết quả JSON và điền workbook thí nghiệm.

Nhiệm vụ: lưu kết quả từng lần chạy ra JSON, rồi điền vào experiments.xlsx từ mẫu
templates/experiment_table_template.xlsx (đừng gõ tay hàng chục dòng, rất dễ sai).

Tên cột của sheet "Experiments" (giữ nguyên, đúng thứ tự mẫu):
    exp_id, group, description, loss, optimizer, lr, weight_decay, batch, epochs, hidden, dropout,
    clip_norm, precision, init, seed, step0_loss, best_val_loss, best_epoch, final_train_loss,
    final_val_loss, val_acc, val_macro_f1, time_per_epoch_s, peak_mem_MB, diverged,
    eval_acc, eval_macro_f1, figure_file, notes
(các cột công thức ở cuối bảng mẫu tự tính, đừng ghi đè)
"""
from __future__ import annotations

import json
import math
from pathlib import Path


def save_result(result: dict, results_dir: str = "../results") -> str:
    """Ghi result["cfg"], result["history"], result["summary"] (KHÔNG ghi best_state) ra
    <results_dir>/<exp_id>.json. Trả về đường dẫn file. Tạo thư mục nếu chưa có."""
    results_dir = Path(results_dir)
    results_dir.mkdir(parents=True, exist_ok=True)
    payload = {key: result[key] for key in ("cfg", "history", "summary")}
    path = results_dir / f"{result['cfg']['exp_id']}.json"
    with path.open("w", encoding="utf-8") as handle:
        json.dump(_json_safe(payload), handle, ensure_ascii=False, indent=2, allow_nan=False)
    return str(path)


def _json_safe(value):
    if isinstance(value, dict):
        return {key: _json_safe(item) for key, item in value.items()}
    if isinstance(value, (list, tuple)):
        return [_json_safe(item) for item in value]
    if isinstance(value, float) and not math.isfinite(value):
        return None
    if hasattr(value, "item"):
        return _json_safe(value.item())
    return value


def load_results(results_dir: str = "../results") -> list[dict]:
    """Đọc mọi file *.json trong results_dir, trả về danh sách dict (sắp theo exp_id)."""
    results_dir = Path(results_dir)
    if not results_dir.exists():
        return []
    results = []
    for path in sorted(results_dir.glob("*.json")):
        with path.open(encoding="utf-8") as handle:
            result = json.load(handle)
        results.append(result)
    return sorted(results, key=lambda result: result["cfg"]["exp_id"])


def to_row(result: dict, eval_scores: dict | None = None, notes: str = "") -> dict:
    """Biến một kết quả thành một dòng của bảng: gộp cfg + summary (+ eval_acc, eval_macro_f1 nếu có)
    + figure_file = f"figures/{exp_id}.png". Khoá phải trùng tên cột ở đầu file.
    Chỉ truyền eval_scores cho baseline và cấu hình cuối cùng."""
    cfg, summary = result["cfg"], result["summary"]
    accuracy = macro_f1 = None
    if eval_scores is not None:
        accuracy = eval_scores.get("accuracy", eval_scores.get("eval_acc"))
        macro_f1 = eval_scores.get("macro_f1", eval_scores.get("eval_macro_f1"))
    return {
        "exp_id": cfg["exp_id"], "group": cfg.get("group", ""),
        "description": cfg.get("description", ""), "loss": cfg["loss"],
        "optimizer": cfg["optimizer"], "lr": cfg["lr"],
        "weight_decay": cfg.get("weight_decay", 0.0), "batch": cfg["batch"],
        "epochs": cfg["epochs"], "hidden": "-".join(map(str, cfg["hidden"])),
        "dropout": cfg["dropout"],
        "clip_norm": cfg.get("clip_norm") if cfg.get("clip_norm") is not None else "none",
        "precision": cfg["precision"], "init": cfg["init"], "seed": cfg["seed"],
        "step0_loss": summary["step0_loss"], "best_val_loss": summary["best_val_loss"],
        "best_epoch": summary["best_epoch"], "final_train_loss": summary["final_train_loss"],
        "final_val_loss": summary["final_val_loss"], "val_acc": summary["val_acc"],
        "val_macro_f1": summary["val_macro_f1"],
        "time_per_epoch_s": summary["time_per_epoch_s"],
        "peak_mem_MB": summary["peak_mem_MB"], "diverged": summary["diverged"],
        "eval_acc": accuracy, "eval_macro_f1": macro_f1,
        "figure_file": f"figures/{cfg['exp_id']}.png", "notes": notes,
    }


def write_xlsx(rows: list[dict], template_path: str, out_path: str) -> None:
    """Điền các dòng vào sheet "Experiments" của mẫu, từ dòng 2 trở xuống, rồi lưu thành out_path.

    Các bước (openpyxl):
      1. wb = openpyxl.load_workbook(template_path)   # KHÔNG dùng data_only=True (sẽ mất công thức)
      2. ws = wb["Experiments"]; đọc tiêu đề dòng 1 để biết cột nào ứng với khoá nào
      3. với mỗi row: ghi giá trị vào đúng cột; BỎ QUA các cột công thức (step0_gap_vs_lnC, gap_val_minus_train,
         delta_val_f1_vs_base, beyond_noise)
      4. wb.save(out_path)
    Sau khi lưu, mở file bằng Excel/LibreOffice để các công thức tính lại.
    """
    try:
        from openpyxl import load_workbook
    except ImportError as error:
        raise RuntimeError("Cần openpyxl để tạo experiments.xlsx; Colab/Kaggle thường đã cài sẵn.") from error

    template_path, out_path = Path(template_path), Path(out_path)
    if not template_path.is_file():
        raise FileNotFoundError(f"Không tìm thấy mẫu bảng tính: {template_path}")
    out_path.parent.mkdir(parents=True, exist_ok=True)
    workbook = load_workbook(template_path)
    if "Experiments" not in workbook.sheetnames or "Seeds" not in workbook.sheetnames or "Summary" not in workbook.sheetnames:
        raise ValueError("Mẫu cần có các sheet Experiments, Seeds và Summary")

    ws = workbook["Experiments"]
    headers = [cell.value for cell in ws[1]]
    if not headers or headers[0] != "exp_id":
        raise ValueError("Không nhận ra header của sheet Experiments")
    formula_columns = {
        "step0_gap_vs_lnC", "gap_val_minus_train", "delta_val_f1_vs_base", "beyond_noise"
    }
    if len(rows) > ws.max_row - 1:
        raise ValueError(f"Mẫu chỉ dự trù {ws.max_row - 1} thí nghiệm; nhận {len(rows)}")

    for row_index, row in enumerate(rows, start=2):
        for column_index, header in enumerate(headers, start=1):
            if header in formula_columns:
                continue
            ws.cell(row=row_index, column=column_index).value = row.get(header)

    seeds = workbook["Seeds"]
    if not str(seeds.cell(row=1, column=1).value).startswith("exp_id"):
        raise ValueError("Không nhận ra cột nhập exp_id của sheet Seeds")
    baseline_rows = [row for row in rows if row.get("group") == "baseline"]
    if len(baseline_rows) > seeds.max_row - 1:
        raise ValueError("Mẫu Seeds không đủ dòng cho số seed baseline")
    for row_index in range(2, seeds.max_row + 1):
        seeds.cell(row=row_index, column=1).value = None
    for row_index, row in enumerate(baseline_rows, start=2):
        seeds.cell(row=row_index, column=1).value = row["exp_id"]

    summary = workbook["Summary"]
    summary_headers = [cell.value for cell in summary[1]]
    if len(summary_headers) < 8 or summary_headers[0] != "group":
        raise ValueError("Summary sheet needs its original eight columns")
    template_groups = {
        summary.cell(row=row_index, column=1).value
        for row_index in range(2, summary.max_row + 1)
    }
    missing_groups = {row.get("group") for row in rows if row.get("group")} - template_groups
    if missing_groups:
        raise ValueError(f"Mẫu Summary thiếu nhóm: {sorted(missing_groups)}")

    if hasattr(workbook, "calculation"):
        workbook.calculation.fullCalcOnLoad = True
        workbook.calculation.forceFullCalc = True
    workbook.save(out_path)
